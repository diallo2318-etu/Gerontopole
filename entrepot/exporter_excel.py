"""Exporte tout l'entrepôt dans un seul fichier Excel, pour le consulter ou le montrer.

Un onglet « Sommaire » décrit chaque table, puis un onglet par table.
Le fichier est écrit dans DONNEES_ICOPE/entrepot_icope.xlsx (hors dépôt git) :
il contient des données de patients.

Utilisation :  python3 entrepot/exporter_excel.py
(après python3 entrepot/construire_entrepot.py)
"""
import duckdb
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from construire_entrepot import DONNEES, ENTREPOT

SORTIE = DONNEES / "entrepot_icope.xlsx"

DESCRIPTIONS = {
    "icope_depistage": "ICOPE étape 1 : une ligne par dépistage (âge, sexe, commune, modalité, 9 alertes, test retenu, essai probable, one-shot de professionnel)",
    "icope_alerte": "ICOPE étape 1 : les mêmes alertes en format long, une ligne par dépistage et par capacité",
    "icope_alerte_statut": "ICOPE : une ligne par dépistage en alerte, avec le statut de l'alerte (traitée, en cours, non traitée)",
    "icope_traitement": "ICOPE : une ligne par action menée sur une alerte (appel, levée de doute, orientation...) ; une alerte peut en avoir plusieurs",
    "v_alerte_issue": "Vue : une ligne par alerte, avec ce qui lui est arrivé (pas prise en charge, orientée vers un professionnel, déjà suivie, fausse alerte...)",
    "icope_evaluation": "ICOPE étape 2 : une ligne par évaluation approfondie",
    "icope_professionnel": "Professionnels inscrits sur ICOPE",
    "v_parcours_participant": "Vue : une ligne par personne (tests retenus, one-shot, modalité du premier et du dernier test, étape 2)",
    "offre_rpps": "Annuaire Santé : professionnels de santé de Normandie et leurs lieux d'exercice",
    "offre_finess": "FINESS : établissements sanitaires et médico-sociaux de Normandie",
    "offre_mon_bilan_prevention": "Professionnels qui font le bilan de prévention en Normandie",
    "v_offre_par_domaine": "Vue : l'offre utile à chaque capacité ICOPE, d'après l'organigramme",
    "ref_alerte_professionnel": "Métiers et structures par capacité (organigramme du Gérontopôle)",
    "ref_commune": "Les 2 644 communes de Normandie : codes, intercommunalité, population, coordonnées",
    "obs_fragilite_commune": "Observatoire des fragilités : indicateurs par commune",
    "obs_fragilite_epci": "Observatoire des fragilités : indicateurs par intercommunalité",
    "obs_fragilite_departement": "Observatoire des fragilités : indicateurs par département",
    "obs_fragilite_dictionnaire": "Observatoire des fragilités : signification de chaque colonne",
    "obs_bienvieillir_commune": "Observatoire « Bien vieillir » : indicateurs par commune",
    "obs_bienvieillir_epci": "Observatoire « Bien vieillir » : indicateurs par intercommunalité",
    "obs_bienvieillir_departement": "Observatoire « Bien vieillir » : indicateurs par département",
    "obs_bienvieillir_dictionnaire": "Observatoire « Bien vieillir » : signification de chaque colonne",
}
ORDRE = list(DESCRIPTIONS)
ENTETE = PatternFill("solid", fgColor="0F4C5C")


def mettre_en_forme(ws, df):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = ENTETE
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        largeur = max([len(str(col)) + 3] + [len(str(v)) for v in df[col].head(200)])  # +3 : bouton de filtre
        ws.column_dimensions[get_column_letter(i)].width = min(max(largeur + 2, 10), 95)


def oui_non(df):
    """Colonnes vrai/faux écrites « oui » / « non » : LibreOffice affiche sinon tantôt VRAI/FAUX, tantôt 1/0."""
    for col in df.columns:
        valeurs = df[col].dropna()
        if len(valeurs) and valeurs.map(lambda v: isinstance(v, (bool, np.bool_))).all():
            df[col] = df[col].map(lambda v: "oui" if v is True or v is np.True_ else ("non" if v is False or v is np.False_ else None))
    return df


def main():
    con = duckdb.connect(str(ENTREPOT), read_only=True)
    tables = [t for (t,) in con.execute("SELECT table_name FROM information_schema.tables").fetchall()]
    tables = [t for t in ORDRE if t in tables] + sorted(t for t in tables if t not in ORDRE)

    sommaire = pd.DataFrame([{
        "onglet": t,
        "contenu": DESCRIPTIONS.get(t, ""),
        "lignes": con.execute(f'SELECT count(*) FROM "{t}"').fetchone()[0],
        "colonnes": len(con.execute(f'SELECT * FROM "{t}" LIMIT 0').description),
    } for t in tables])

    with pd.ExcelWriter(SORTIE, engine="openpyxl") as xl:
        sommaire.to_excel(xl, sheet_name="Sommaire", index=False)
        mettre_en_forme(xl.sheets["Sommaire"], sommaire)
        for t in tables:
            df = oui_non(con.execute(f'SELECT * FROM "{t}"').df())
            df.to_excel(xl, sheet_name=t[:31], index=False)
            mettre_en_forme(xl.sheets[t[:31]], df)
            print(f"  {t:32s} {len(df):>7} lignes")
    con.close()
    print(f"\nFichier écrit : {SORTIE}")


if __name__ == "__main__":
    main()
